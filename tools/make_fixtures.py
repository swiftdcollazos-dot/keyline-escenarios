"""Genera tests/fixtures/excel_cases.json: para cada variante cambia datos de la
hoja 1, recalcula el libro con LibreOffice y guarda los resultados de Excel.

Uso: python3 tools/make_fixtures.py  (necesita openpyxl y LibreOffice)
"""
import json, os, shutil, subprocess, sys, tempfile
import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, "tests", "fixtures", "La_Esperanza_calculos_paso_a_paso.xlsx")
OUT = os.path.join(ROOT, "tests", "fixtures", "excel_cases.json")

VARIANTES = {
    "referencia": {},
    "evaporacion_7": {"Evap_dia": 7, "Infil_dia": 1.5},
    "reses_80": {"N_reses": 80, "Consumo_res": 60},
    "hondas": {"Hmax_occ": 5, "Hmax_ori": 5, "Espejo_occ": 3760, "Espejo_ori": 4100},
    "canal_arenoso": {"Pend_canal": 200, "Suelo_canal": "Franco arenoso", "b_canal": 0.4, "y_canal": 0.35},
    "manguera_20": {"D_tubo": 20, "Q_casa": 0.3, "F_trazado": 1.4, "Cota_casa": 382},
    "puntos_clave": {"PK_occ": 388.6, "PK_ori": 386.2, "Int_franjas": 15, "Margen_grav": 1.5},
    "potreros": {"N_potreros": 16, "Descanso": 42, "Ancho_camino": 5, "Ronda": 20, "Area_constr": 0.5, "Dist_cacao": 3.5},
    "escala_y_curvas": {"Barra_px": 82, "Int_curvas": 0.5, "Lluvia_anual": 2000, "C_anual": 0.25},
    "muros": {"Borde_libre": 1.2, "Corona": 4, "Talud_muro": 2.5, "Lam_vert": 0.4, "Coef_vert": 1.6, "C_tormenta": 0.6, "I_diseno": 120},
    "verano_150": {"Dias_verano": 150},
}

def recalc(path, outdir):
    subprocess.run(["soffice", "--headless", "--calc", "--convert-to", "xlsx", "--outdir", outdir, path],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=180)
    return os.path.join(outdir, os.path.basename(path))

def valores(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    res = {}
    for name, dn in wb.defined_names.items():
        for sheet, ref in dn.destinations:
            if ":" in ref:
                continue
            v = wb[sheet][ref.replace("$", "")].value
            if isinstance(v, (int, float, str)):
                res[name] = v
    extra = {
        "4 Balance de verano": ["E21", "F21", "E22", "F22", "E24", "E37", "F37", "E38"] + [f"{c}{r}" for r in range(28, 34) for c in "CEFG"],
        "5 Simulación 120 días": ["G14", "H14", "G15", "H15"],
        "6 Muros y vertederos": ["E10", "F10", "E11", "F11", "E18", "F18", "E19", "F19", "E20", "F20", "E21", "F21"],
        "7 Canales 1 en 300": ["E11", "F11", "E12", "E15", "F15", "F18", "E41", "E42", "E43", "E44", "E45", "E47", "E48", "E49", "F49", "E50"],
        "8 Casa y bebederos": ["E10", "E11", "E15", "E20", "E24", "E25", "E26", "G37", "G38", "G39", "G40", "H37", "H38", "H39", "H40", "I37", "I40"],
        "9 Árboles cacao ronda": ["E10", "E12", "E44", "E45", "E46"],
        "10 Potreros": ["E8", "E12", "E18", "E19", "E21", "E22", "E23", "G27", "G28", "G29"] + [f"H{r}" for r in range(33, 45)],
        "11 Verificación": [f"E{r}" for r in range(7, 17)] + [f"G{r}" for r in range(7, 17)],
    }
    for sheet, cells in extra.items():
        for c in cells:
            v = wb[sheet][c].value
            if isinstance(v, (int, float, str)):
                res[f"{sheet.split(' ')[0]}!{c}"] = v
    return res

def main():
    base = openpyxl.load_workbook(SRC)
    filas = {base["1 Datos"].cell(row=r, column=2).value: r for r in range(1, 80)}
    casos = []
    with tempfile.TemporaryDirectory() as tmp:
        for nombre, cambios in VARIANTES.items():
            wb = openpyxl.load_workbook(SRC)
            for k, v in cambios.items():
                wb["1 Datos"].cell(row=filas[k], column=4).value = v
            src = os.path.join(tmp, "in", nombre + ".xlsx")
            os.makedirs(os.path.dirname(src), exist_ok=True)
            wb.save(src)
            out = recalc(src, os.path.join(tmp, "out"))
            casos.append({"nombre": nombre, "cambios": cambios, "excel": valores(out)})
            print(nombre, "ok", file=sys.stderr)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(casos, f, ensure_ascii=False, indent=1)

if __name__ == "__main__":
    main()
